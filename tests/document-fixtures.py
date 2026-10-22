"""Office/PDF fixtures with real editable structure, used by runtime and desktop tests."""
from pathlib import Path
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches as SlideInches
from reportlab.pdfgen import canvas
from PIL import Image


def create(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    image = folder / 'mark.png'
    Image.new('RGB', (160, 80), '#327d92').save(image)
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = 'KOAN / Delivery brief'
    doc.add_heading('Document acceptance', 0)
    doc.add_paragraph('An editable brief with a table, image and a second page.')
    table = doc.add_table(rows=3, cols=2)
    table.style = 'Light Shading Accent 1'
    for row, values in zip(table.rows, [('Project', 'Hours'), ('Video Dancer', '4'), ('Trismegistus', '3')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.add_picture(str(image), width=Inches(2))
    doc.add_page_break()
    doc.add_heading('Next steps', 1)
    doc.add_paragraph('Revise the estimate and preserve the source file.')
    doc.save(folder / 'brief.docx')
    book = Workbook()
    sheet = book.active
    sheet.title = 'Estimate'
    for row in [('Project', 'Hours', 'Rate', 'Value'), ('Video Dancer', 4, 100, '=B2*C2'), ('Trismegistus', 3, 120, '=B3*C3'), ('Total', '=SUM(B2:B3)', None, '=SUM(D2:D3)')]:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='234559')
    sheet.column_dimensions['A'].width = 22
    for key in ['B', 'C', 'D']:
        sheet.column_dimensions[key].width = 15
    chart = BarChart()
    chart.title = 'Value by project'
    chart.add_data(Reference(sheet, min_col=4, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, 'A7')
    sheet.print_area = 'A1:H23'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    book.save(folder / 'estimate.xlsx')
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[0])
    slide.shapes.title.text = 'KOAN delivery'
    slide.placeholders[1].text = 'Editable project summary'
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = 'Project value'
    data = CategoryChartData()
    data.categories = ['Video Dancer', 'Trismegistus']
    data.add_series('Value', (400, 360))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, SlideInches(1), SlideInches(1.7), SlideInches(8), SlideInches(4.5), data)
    slides.save(folder / 'summary.pptx')
    pdf = canvas.Canvas(str(folder / 'form.pdf'))
    pdf.setFont('Helvetica-Bold', 20)
    pdf.drawString(60, 760, 'KOAN approval')
    pdf.setFont('Helvetica', 12)
    pdf.drawString(60, 710, 'Project')
    pdf.acroForm.textfield(name='project', x=120, y=695, width=250, height=25, value='Video Dancer')
    pdf.showPage()
    pdf.drawString(60, 760, 'Delivery notes')
    pdf.save()


if __name__ == '__main__':
    import sys
    create(sys.argv[1])
