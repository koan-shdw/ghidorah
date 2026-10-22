import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP / 'backend'))
import documents
spec = importlib.util.spec_from_file_location('document_fixtures', APP / 'tests/document-fixtures.py')
fixtures = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixtures)


class Documents(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.folder = Path(tempfile.mkdtemp(prefix='documents-runtime-', dir=APP / 'work'))
        fixtures.create(cls.folder)
        print('Document evidence: ' + str(cls.folder), flush=True)

    def render(self, file, target, recalculate=False):
        return documents.render_request({'root': str(self.folder), 'source': str(self.folder / file),
                                         'destination': str(self.folder / target), 'recalculate': recalculate})

    def test_create_edit_render_and_recalculate(self):
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        from pypdf import PdfReader, PdfWriter
        rendered = {}
        for file, count in [('brief.docx', 2), ('estimate.xlsx', 1), ('summary.pptx', 2), ('form.pdf', 2)]:
            before = (self.folder / file).read_bytes()
            if not file.endswith('.pdf'):
                self.assertIn('KOAN' if file.endswith('.pptx') else 'Video Dancer', str(documents.inspect_office(self.folder / file, file)))
            value = self.render(file, file + '-render')
            rendered[file] = value
            self.assertEqual(value['pageCount'], count)
            self.assertTrue(all(Path(p).stat().st_size > 1000 for p in value['pages']))
            self.assertEqual((self.folder / file).read_bytes(), before)
        doc = Document(self.folder / 'brief.docx')
        doc.tables[0].cell(1, 1).text = '5'
        doc.save(self.folder / 'revised.docx')
        self.assertEqual(len(Document(self.folder / 'revised.docx').inline_shapes), 1)
        self.assertIn('Delivery brief', doc.sections[0].header.paragraphs[0].text)
        self.render('revised.docx', 'revised-render')
        book = load_workbook(self.folder / 'estimate.xlsx')
        book.active['B2'] = 5
        book.save(self.folder / 'changed.xlsx')
        result = self.render('changed.xlsx', 'recalculated', True)
        values = load_workbook(result['file'], data_only=True)
        self.assertEqual(values.active['D4'].value, 860)
        self.assertEqual(values.active['B4'].value, 8)
        self.assertEqual(result['formulaErrors'], [])
        formulas = load_workbook(result['file'], data_only=False)
        self.assertEqual(formulas.active['D2'].value, '=B2*C2')
        self.assertEqual(len(formulas.active._charts), 1)
        self.assertTrue(formulas.active['A1'].font.bold)
        slides = Presentation(self.folder / 'summary.pptx')
        slides.slides[0].shapes.title.text = 'Revised KOAN delivery'
        slides.save(self.folder / 'revised.pptx')
        self.assertEqual(len(slides.slides), 2)
        self.assertTrue(any(shape.has_chart for shape in slides.slides[1].shapes))
        self.render('revised.pptx', 'revised-slides')
        writer = PdfWriter(clone_from=str(self.folder / 'form.pdf'))
        writer.update_page_form_field_values(writer.pages[0], {'project': 'Trismegistus'}, auto_regenerate=False)
        writer.write(self.folder / 'filled.pdf')
        self.assertEqual(PdfReader(self.folder / 'filled.pdf').get_fields()['project']['/V'], 'Trismegistus')
        self.render('filled.pdf', 'filled-render')
        (self.folder / 'verified.json').write_text(json.dumps(rendered, indent=2))

    def test_bounds_active_content_external_data_and_overwrite(self):
        for addition in [('word/vbaProject.bin', b'x'), ('word/embeddings/object.bin', b'x'),
                         ('word/_rels/unsafe.rels', b'<Relationships><Relationship TargetMode="External" Type="image" Target="https://example.com/pixel"/></Relationships>')]:
            target = self.folder / ('unsafe-' + str(len(addition[0])) + '.docx')
            target.write_bytes((self.folder / 'brief.docx').read_bytes())
            with zipfile.ZipFile(target, 'a') as archive:
                archive.writestr(*addition)
            with self.assertRaises(ValueError):
                documents.office_check(target, target.name)
        with self.assertRaises(ValueError):
            documents.render_request({'root': str(self.folder), 'source': str(APP / 'package.json'), 'destination': str(self.folder / 'outside')})
        with self.assertRaises((ValueError, FileExistsError)):
            self.render('brief.docx', 'brief.docx')
        self.assertEqual(documents.office_check(self.folder / 'summary.pptx', 'summary.pptx'), '.pptx')
        from openpyxl import load_workbook
        book = load_workbook(self.folder / 'estimate.xlsx')
        book.active['F1'] = '=WEBSERVICE("https://example.com/data")'
        book.save(self.folder / 'external.xlsx')
        with self.assertRaisesRegex(ValueError, 'External data'):
            documents.office_check(self.folder / 'external.xlsx', 'external.xlsx')


if __name__ == '__main__':
    unittest.main()
