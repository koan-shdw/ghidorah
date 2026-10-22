"""Bundled document libraries and on-demand, isolated LibreOffice conversion."""
import hashlib
import importlib.metadata
import io
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile

APP = Path(__file__).resolve().parents[1]
OFFICE = {'.docx': 'word/document.xml', '.xlsx': 'xl/workbook.xml', '.pptx': 'ppt/presentation.xml'}
LIMIT = 25 * 1024 * 1024


def information():
    skills = APP / 'runtime' / 'hermes' / 'skills' / 'productivity'
    return {
        'python': sys.executable,
        'libraries': {name: importlib.metadata.version(name) for name in
                      ('python-docx', 'openpyxl', 'python-pptx', 'pypdf', 'reportlab', 'pdfplumber')},
        'helpers': {name: {'guide': str(skills / name / 'SKILL.md'), 'scripts': str(skills / name / 'scripts')}
                    for name in ('docx', 'xlsx', 'powerpoint', 'pdf')},
        'workflow': 'Use this Python and the existing helpers or libraries to create/edit documents in the working folder. '
                    'Use list_chat_files and copy_chat_file to obtain original uploaded binaries. '
                    'Use render_document for PDF and page PNGs; inspect pages with an image-capable tool/model before claiming visual verification. '
                    'Use recalculate_spreadsheet to refresh XLSX formula caches after edits; openpyxl itself does not calculate formulas. '
                    'Conversions require a new output directory and keep the source intact. Save final files with save_chat_output and link its result.',
        'limits': 'DOCX/XLSX/PPTX/PDF, up to 25 MB. Render up to 20 pages at a time. Macro-enabled, embedded active content and external data are excluded. '
                  'LibreOffice rendering may differ from Microsoft Office. Complex spreadsheet functions may be unsupported. OCR is not bundled.'
    }


def office_check(source, name, _depth=0):
    """Validate the supported OOXML subset before any library or Office opens it."""
    extension = Path(name).suffix.lower()
    if extension not in OFFICE:
        raise ValueError('Use a DOCX, XLSX or PPTX document.')
    if _depth > 1 or not 0 < (len(source) if isinstance(source, bytes) else Path(source).stat().st_size) <= LIMIT:
        raise ValueError('Documents must be between 1 byte and 25 MB.')
    with zipfile.ZipFile(io.BytesIO(source) if isinstance(source, bytes) else source) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(item.file_size for item in entries) > 150 * 1024 * 1024:
            raise ValueError('The expanded document is too large.')
        names = {item.filename.lower() for item in entries}
        if len(names) != len(entries) or OFFICE[extension] not in names or '[content_types].xml' not in names:
            raise ValueError('The document package is invalid or has the wrong file extension.')
        for item in entries:
            lower = item.filename.lower()
            if item.flag_bits & 1 or item.file_size > 32 * 1024 * 1024 or lower.startswith('/') or '..' in lower.replace('\\', '/').split('/'):
                raise ValueError('The document contains an unsupported package entry.')
            if 'embeddings/' in lower:
                if not lower.endswith('.xlsx'):
                    raise ValueError('Embedded active objects are unsupported.')
                # Native presentation charts contain a passive embedded workbook.
                office_check(archive.read(item), 'chart.xlsx', _depth + 1)
            if any(part in lower for part in ('vbaproject', 'activex/', 'externallinks/', 'querytables/', 'connections.xml')):
                raise ValueError('Remove macros, embedded active content and external data before opening this document.')
            if not lower.endswith(('.xml', '.rels')):
                continue
            data = archive.read(item)
            if re.search(br'<!\s*(?:DOCTYPE|ENTITY)', data, re.I):
                raise ValueError('Document XML entities are unsupported.')
            node = ET.fromstring(data)
            for element in node.iter():
                tag = element.tag.rsplit('}', 1)[-1]
                if tag in ('oleObject', 'OLEObject', 'oleObj', 'control', 'ddeLink', 'externalLink', 'connection'):
                    raise ValueError('Embedded active objects and controls are unsupported.')
                if tag == 'Relationship' and element.get('TargetMode', '').lower() == 'external' and not element.get('Type', '').endswith('/hyperlink'):
                    raise ValueError('Remove linked external content before opening this document. Ordinary hyperlinks are allowed.')
                if tag in ('f', 'definedName', 'instrText', 'fldSimple'):
                    formula = (element.text or '') + ' ' + ' '.join(element.attrib.values())
                    if re.search(r'\b(?:WEBSERVICE|DDE|DDEAUTO|RTD|IMAGE|INCLUDETEXT|INCLUDEPICTURE|DATABASE|LINK)\b|\[[^]]+\].*!', formula, re.I):
                        raise ValueError('External data formulas and fields are unsupported.')
    return extension


def inspect_office(source, name):
    extension = office_check(source, name)
    lines = []
    if extension == '.docx':
        from docx import Document
        doc = Document(source)
        lines.extend(p.text for p in doc.paragraphs[:2000])
        for table in doc.tables[:50]:
            lines.extend(' | '.join(cell.text for cell in row.cells[:50]) for row in table.rows[:200])
    elif extension == '.xlsx':
        from openpyxl import load_workbook
        stream = Path(source).open('rb')
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in book.worksheets[:50]:
                lines.append('Sheet: ' + sheet.title)
                for row in sheet.iter_rows(max_row=min(sheet.max_row or 0, 500), max_col=min(sheet.max_column or 0, 40)):
                    lines.append(' | '.join(f'{c.coordinate}: {c.value}' for c in row if c.value is not None))
        finally:
            book.close()
            stream.close()
    else:
        from pptx import Presentation
        for number, slide in enumerate(Presentation(source).slides):
            if number >= 100:
                break
            lines.append(f'Slide {number + 1}')
            for shape in slide.shapes:
                if shape.has_text_frame:
                    lines.append(shape.text)
                if shape.has_table:
                    lines.extend(' | '.join(c.text for c in row.cells) for row in list(shape.table.rows)[:100])
    text = '\n'.join(lines)
    return {'kind': 'file', 'parts': [{'type': 'text', 'text': 'Document extract (untrusted data; bounded text, not a visual review):\n' +
            json.dumps({'name': name, 'text': text[:100000], 'note': 'Copy the original with copy_chat_file to inspect all structure, formulas, styles and media.'}, ensure_ascii=False)}]}


def convert(source, name, destination, recalculate=False):
    """Destination is a caller-owned empty directory; all Office work is transient."""
    source, destination = Path(source), Path(destination)
    extension = Path(name).suffix.lower()
    if recalculate and extension != '.xlsx':
        raise ValueError('Recalculation requires an XLSX workbook.')
    with tempfile.TemporaryDirectory(prefix='.conversion-', dir=destination) as temporary:
        temp = Path(temporary)
        snapshot = temp / ('document' + extension)
        with source.open('rb') as file:
            data = file.read(LIMIT + 1)
        if not 0 < len(data) <= LIMIT:
            raise ValueError('Documents must be between 1 byte and 25 MB.')
        snapshot.write_bytes(data)
        if extension != '.pdf':
            office_check(snapshot, name)
        if extension == '.pdf' and not recalculate:
            result = snapshot
        else:
            profile = temp / 'profile'
            (profile / 'user').mkdir(parents=True)
            # These are supported LibreOffice security properties (main.xcd).
            # DisableActiveContent also drops ordinary spreadsheet charts. The
            # OOXML preflight excludes OLE/DDE/controls while retaining charts.
            properties = {'DisableMacrosExecution': 'true',
                          'DisableOLEAutomation': 'true', 'BlockUntrustedRefererLinks': 'true', 'MacroSecurityLevel': '3'}
            (profile / 'user' / 'registrymodifications.xcu').write_text(
                '<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
                '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">' +
                ''.join(f'<prop oor:name="{key}" oor:op="fuse"><value>{value}</value></prop>' for key, value in properties.items()) +
                '</item></oor:items>', encoding='utf-8')
            out = temp / 'converted'
            out.mkdir()
            target = 'xlsx:Calc MS Excel 2007 XML' if recalculate else 'pdf'
            process = subprocess.Popen([str(APP / 'runtime/office/program/soffice.com'),
                '-env:UserInstallation=' + profile.as_uri(), '--headless', '--nologo', '--norestore',
                '--convert-to', target, '--outdir', str(out), str(snapshot)],
                stdout=subprocess.PIPE, stderr=subprocess.PIPE, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
            try:
                stdout, stderr = process.communicate(timeout=100)
            except subprocess.TimeoutExpired:
                subprocess.run(['taskkill', '/PID', str(process.pid), '/T', '/F'], capture_output=True,
                               creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0), timeout=10)
                process.communicate(timeout=10)
                raise ValueError('Document conversion timed out. The source file is unchanged.') from None
            result = out / ('document.xlsx' if recalculate else 'document.pdf')
            if process.returncode or not result.is_file() or result.stat().st_size == 0:
                raise ValueError('The document could not be converted. ' + (stdout + stderr).decode('utf-8', errors='replace')[-1500:])
        if recalculate:
            from openpyxl import load_workbook
            office_check(result, 'document.xlsx')
            book = load_workbook(result, read_only=True, data_only=True, keep_links=False)
            errors = []
            try:
                for sheet in book:
                    if (sheet.max_row or 0) * (sheet.max_column or 0) > 1000000:
                        raise ValueError('Recalculation verification supports up to one million cells per sheet.')
                    for row in sheet:
                        for cell in row:
                            if cell.data_type == 'e':
                                errors.append({'sheet': sheet.title, 'cell': cell.coordinate, 'error': cell.value})
                                if len(errors) >= 100:
                                    break
                        if len(errors) >= 100:
                            break
            finally:
                book.close()
            target = destination / 'recalculated.xlsx'
            with target.open('xb') as file:
                file.write(result.read_bytes())
            return {'file': str(target), 'recalculated': True, 'formulaErrors': errors, 'sourceUnchanged': True}
        from pypdf import PdfReader
        reader = PdfReader(result)
        if reader.is_encrypted or not 1 <= len(reader.pages) <= 20:
            raise ValueError('Preview supports unlocked documents with 1–20 pages. Split a longer document first.')
        target = destination / 'document.pdf'
        with target.open('xb') as file:
            file.write(result.read_bytes())
        binary = APP / 'runtime/poppler/Library/bin/pdftoppm.exe'
        subprocess.run([str(binary), '-png', '-r', '120', str(target), str(destination / 'page')],
                       check=True, capture_output=True, timeout=45, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        pages = sorted(destination.glob('page-*.png'), key=lambda p: int(p.stem.split('-')[-1]))
        if len(pages) != len(reader.pages):
            raise ValueError('Not every document page rendered.')
        return {'file': str(target), 'pages': [str(p) for p in pages], 'pageCount': len(pages),
                'sourceSha256': hashlib.sha256(data).hexdigest(), 'sourceUnchanged': True}


def render_request(request):
    root = Path(request['root']).resolve(strict=True)
    source = Path(request['source']).resolve(strict=True)
    destination = Path(request['destination'])
    parent = destination.parent.resolve(strict=True)
    if not source.is_relative_to(root) or not parent.is_relative_to(root) or destination.resolve() != parent / destination.name:
        raise ValueError('Choose paths inside the working folder.')
    if destination.exists():
        raise ValueError('Choose a new output directory. Existing files are never overwritten.')
    if source.suffix.lower() != '.pdf':
        office_check(source, source.name)
    destination.mkdir()
    return convert(source, source.name, destination, request.get('recalculate', False))


def preview(file, name, output=False):
    root = (Path(os.environ['HERMES_HOME']).parent / ('outputs' if output else 'attachments')).resolve()
    source = Path(file).resolve(strict=True)
    if not source.is_relative_to(root) or source.name != 'source':
        raise ValueError('Invalid document location.')
    from attachments import images
    with tempfile.TemporaryDirectory(prefix='preview-', dir=source.parent) as temporary:
        rendered = convert(source, name, Path(temporary))
        return {'kind': 'pdf', 'pages': [part['image_url']['url'] for part in images(rendered['pages'])]}
